import io,json,hashlib
from datetime import date
from types import SimpleNamespace
from clearspace.analysis import experience_timeline,verify_assessment
from clearspace.privacy import redact
from clearspace.files import extract

def nlp(text):
    name='Ali Khan';start=text.find(name)
    return SimpleNamespace(ents=[SimpleNamespace(label_='PERSON',start_char=start,end_char=start+len(name))] if start>=0 else [])

def test_redaction_preserves_work_dates():
    text='Ali Khan\nali@example.com +92 300 1234567\nExperience\nEngineer Jan 2020 - Jan 2025\nWorked with Ali Khan on products.'
    safe=redact(text,nlp,'Ali Khan')
    assert 'Ali Khan' not in safe and 'ali@example.com' not in safe and '1234567' not in safe
    assert 'Jan 2020 - Jan 2025' in safe

def test_overlapping_work_ranges_education_and_year_only():
    text='Experience\nEngineer Jan 2020 - Jan 2024\nLead Jan 2022 - Jan 2025\nEducation\nDegree Jan 2015 - Jan 2019'
    assert experience_timeline(text,date(2026,1,1))['years']==5
    assert experience_timeline('Experience\nEngineer 2020 - 2025')['years'] is None

def test_transferable_partial_and_unverified_evidence():
    row={'requirement':'Redux','category':'skill','importance':'required','status':'met','jobEvidence':'Redux experience','cvEvidence':'Zustand','basis':'transferable'}
    result=verify_assessment({'requirements':[row]},'Built UI with Zustand','Redux experience')
    assert result['score']==50 and result['requirements'][0]['status']=='partial'
    row['cvEvidence']='Invented qualification'
    assert verify_assessment({'requirements':[row]},'Built UI with Zustand','Redux experience')['score']==0

def test_docx_and_invalid_uploads():
    import pytest
    from docx import Document
    doc=Document();doc.add_paragraph('Ali Khan');doc.add_paragraph('Experience');doc.add_paragraph('Built customer software with Python and TypeScript and tested API integrations for production users.')
    stream=io.BytesIO();doc.save(stream)
    assert 'Python' in extract(stream.getvalue(),'cv.docx')
    for content,name in [(b'fake','resume.pdf'),(b'fake','resume.doc'),(b'fake','resume.docx')]:
        with pytest.raises(Exception):extract(content,name)

def configure_app():
    import httpx
    from clearspace.main import app
    app.state.public=True;app.state.secret='x'*48;app.state.token_hashes=[];app.state.key='server-only-secret';app.state.provider='https://api.thehive.ai/api/v3/chat/completions'
    app.state.nlp=nlp;app.state.prompts=json.loads(__import__('pathlib').Path('clearspace/prompts.json').read_text())
    class Redis:
        def __init__(self):self.count={};self.leases={};self.fail=False
        async def eval(self,script,count,key,*args):
            if self.fail:raise RuntimeError('down')
            if "INCR" in script:
                self.count[key]=self.count.get(key,0)+1;return [self.count[key],int(args[0])]
            self.leases.pop(key,None);return 1
        async def set(self,key,owner,**kwargs):
            if key in self.leases:return False
            self.leases[key]=owner;return True
    app.state.redis=Redis();seen=[]
    def provider(request):
        body=json.loads(request.content);seen.append(body)
        assert request.headers['Authorization']=='Bearer server-only-secret'
        content={'requirements':[{'requirement':'Python','category':'skill','importance':'required','status':'met','jobEvidence':'Python','cvEvidence':'Python','basis':'exact'}],'suggestions':[]}
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(content)}}]})
    app.state.http=httpx.AsyncClient(transport=httpx.MockTransport(provider))
    return app,seen

CV='Ali Khan\nali@example.com +92 300 1234567\nExperience\nEngineer Jan 2020 - Jan 2025\nBuilt Python software and tested APIs for real users.'
JOB='We require Python experience for engineering software and integrating APIs. Production software delivery is required.'

def test_provider_receives_redacted_cv_and_errors_do_not_echo_inputs():
    from fastapi.testclient import TestClient
    app,seen=configure_app();client=TestClient(app)
    result=client.post('/api/clearspace/job-match',json={'cv':CV,'job':JOB,'candidateName':'Ali Khan'})
    assert result.status_code==200,result.text
    sent=json.dumps(seen);assert 'ali@example.com' not in sent and 'Ali Khan' not in sent and '1234567' not in sent
    invalid=client.post('/api/clearspace/job-match',json={'cv':'private-secret','job':JOB})
    assert invalid.status_code==422 and 'private-secret' not in invalid.text
    assert result.headers['Cache-Control']=='no-store'

def test_auth_quota_and_redis_fail_closed():
    from fastapi.testclient import TestClient
    app,_=configure_app();client=TestClient(app);app.state.public=False
    assert client.post('/api/clearspace/job-match',json={}).status_code==401
    app.state.public=True;app.state.redis.fail=True
    assert client.post('/api/clearspace/job-match',json={}).status_code==503
    app.state.redis.fail=False
    for _ in range(60):client.post('/api/clearspace/job-match',json={})
    assert client.post('/api/clearspace/job-match',json={}).status_code==429

def test_upload_preview_does_not_call_provider_and_body_limit():
    from fastapi.testclient import TestClient
    app,seen=configure_app();client=TestClient(app)
    result=client.post('/api/clearspace/resume-extract',files={'file':('cv.txt',CV.encode(),'text/plain')})
    assert result.status_code==200,result.text
    assert 'Ali Khan' not in result.json()['text'] and not seen
    result=client.post('/api/clearspace/job-match',content=b'x'*(6*1024*1024+1),headers={'Content-Type':'application/json'})
    assert result.status_code==413
def test_chunked_body_limit_invalid_host_and_bad_token():
    from fastapi.testclient import TestClient
    app,_=configure_app();client=TestClient(app)
    result=client.post('/api/clearspace/job-match',content=iter([b'x'*(4*1024*1024),b'x'*(4*1024*1024)]),headers={'Content-Type':'application/json'})
    assert result.status_code==413,result.text
    assert client.get('/health',headers={'Host':'evil.example'}).status_code==400
    assert client.post('/api/clearspace/job-match',json={},headers={'Authorization':'Bearer untrusted'}).status_code==401
