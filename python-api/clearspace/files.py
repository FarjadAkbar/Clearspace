import io,sys,json,zipfile
from pathlib import Path

def extract(data,filename):
    ext=Path(filename).suffix.lower()
    if ext=='.txt':text=data.decode('utf-8-sig')
    elif ext=='.pdf':
        if not data.startswith(b'%PDF-'):raise ValueError('Invalid PDF')
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=True)
        if reader.is_encrypted or len(reader.pages)>20:raise ValueError('Encrypted PDF or too many pages')
        chunks=[]
        for page in reader.pages:
            chunks.append(page.extract_text() or '')
            if sum(len(c) for c in chunks)>30000:raise ValueError('Resume too long')
        text='\n'.join(chunks)
    elif ext=='.docx':
        if not data.startswith(b'PK'):raise ValueError('Invalid DOCX')
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos=archive.infolist()
            if len(infos)>256 or sum(i.file_size for i in infos)>16*1024*1024:raise ValueError('Document exceeds extraction limits')
            if 'word/document.xml' not in archive.namelist() or any('vbaproject' in i.filename.lower() for i in infos):raise ValueError('Invalid Word document')
            for i in infos:
                if i.filename.endswith('.rels') and b'TargetMode="External"' in archive.read(i):raise ValueError('External document references are not allowed')
        from docx import Document
        doc=Document(io.BytesIO(data));chunks=[p.text for p in doc.paragraphs]
        chunks.extend(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        text='\n'.join(chunks)
    else:raise ValueError('Use PDF, DOCX or UTF-8 TXT. Convert legacy DOC to DOCX first.')
    if not 80<=len(text.strip())<=30000:raise ValueError('Not enough readable text or resume too long. Scanned PDFs need text extraction first.')
    return text.strip()

if __name__=='__main__':
    try:
        if sys.platform!='win32':
            import resource
            resource.setrlimit(resource.RLIMIT_AS,(512*1024*1024,512*1024*1024))
            resource.setrlimit(resource.RLIMIT_CPU,(8,8))
        request=json.loads(sys.stdin.read(8*1024*1024))
        import base64
        print(json.dumps({'text':extract(base64.b64decode(request['data'],validate=True),request['filename'])}))
    except Exception:
        print(json.dumps({'error':'Cannot safely read this file. Upload a text-based PDF, DOCX or TXT.'}));sys.exit(1)
